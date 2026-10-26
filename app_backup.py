import hashlib
import hmac
import os
import re
import sqlite3
import tempfile
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import pdfplumber
import pytesseract
from docx import Document
from dotenv import load_dotenv
from flask import Flask, flash, g, redirect, render_template, request, session, url_for
from PIL import Image, UnidentifiedImageError
from werkzeug.utils import secure_filename

load_dotenv()
BASE_DIR = Path(__file__).resolve().parent
ALLOWED_EXTENSIONS = {"docx", "pdf", "png", "jpg", "jpeg"}
MAX_UPLOAD_BYTES = 12 * 1024 * 1024

app = Flask(__name__)
app.config.update(
    SECRET_KEY=os.getenv("APP_SECRET_KEY") or os.urandom(32).hex(),
    DATABASE=os.getenv("DATABASE_PATH", str(BASE_DIR / "instance" / "certificates.sqlite3")),
    LOCAL_STORAGE=os.getenv("LOCAL_STORAGE", str(BASE_DIR / "instance" / "uploaded_certificates")),
    MAX_CONTENT_LENGTH=MAX_UPLOAD_BYTES,
    TESSERACT_CMD=os.getenv("TESSERACT_CMD", "").strip(),
    DEMO_ACCESS_CODE=os.getenv("DEMO_ACCESS_CODE", ""),
)
if app.config["TESSERACT_CMD"]:
    pytesseract.pytesseract.tesseract_cmd = app.config["TESSERACT_CMD"]

FIELD_PATTERNS = {
    "student_name": [
        r"(?:student\s+name|name\s+of\s+(?:the\s+)?student|awarded\s+to|presented\s+to|certifies\s+that)\s*[:\-]?\s*([A-Z][A-Za-z.' -]{1,80}?)(?=\s+(?:has|for|in|on|with|having|who)\b|[,.;\n]|$)",
    ],
    "certificate_id": [
        r"(?:certificate\s+(?:id|no\.?|number)|credential\s+id)\s*[:#\-]?\s*([A-Za-z0-9][A-Za-z0-9_./-]{2,80})",
    ],
    "course_name": [
        r"(?:course\s+(?:name|title)|program(?:me)?\s+(?:name|title)|completed\s+(?:the\s+)?course|certified\s+in)\s*[:\-]?\s*([^\n,.;]{2,100})",
    ],
    "issue_date": [
        r"(?:issue\s+date|date\s+of\s+issue|issued\s+on|date\s+issued)\s*[:\-]?\s*([A-Za-z0-9 ,./-]{4,40})",
    ],
    "institution": [
        r"(?:issuing\s+institution|institution|issued\s+by|awarded\s+by)\s*[:\-]?\s*([^\n,.;]{2,120})",
    ],
}

def get_db():
    if "db" not in g:
        database_path = Path(app.config["DATABASE"])
        database_path.parent.mkdir(parents=True, exist_ok=True)
        g.db = sqlite3.connect(database_path)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys = ON")
    return g.db

@app.teardown_appcontext
def close_db(_error=None):
    database = g.pop("db", None)
    if database is not None:
        database.close()

def init_db():
    database = get_db()
    database.executescript("""
        CREATE TABLE IF NOT EXISTS certificates (
            id TEXT PRIMARY KEY,
            original_filename TEXT NOT NULL,
            certificate_id TEXT NOT NULL,
            student_name TEXT NOT NULL,
            course_name TEXT NOT NULL,
            issue_date TEXT NOT NULL,
            institution TEXT NOT NULL,
            expected_sha256 TEXT NOT NULL,
            local_path TEXT NOT NULL,
            uploaded_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS verification_events (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            certificate_ref TEXT NOT NULL,
            verified_at TEXT NOT NULL,
            result TEXT NOT NULL,
            FOREIGN KEY (certificate_ref) REFERENCES certificates(id)
        );
    """)
    columns = {row["name"] for row in database.execute("PRAGMA table_info(certificates)").fetchall()}
    if "local_path" not in columns:
        database.execute("ALTER TABLE certificates ADD COLUMN local_path TEXT NOT NULL DEFAULT ''")
    database.commit()

with app.app_context():
    init_db()

def access_code_required():
    return bool(app.config["DEMO_ACCESS_CODE"])

@app.before_request
def enforce_demo_access():
    deployment_requires_code = (
        bool(os.getenv("RENDER")) or os.getenv("FLASK_ENV") == "production"
        or os.getenv("APP_ENV") == "production"
    )
    configured_code = app.config["DEMO_ACCESS_CODE"]
    allowed = {"access", "submit_access_code", "static"}
    if deployment_requires_code and not configured_code and request.endpoint not in allowed:
        return render_template("access.html", configuration_error=True), 503
    if configured_code and not session.get("demo_access_granted") and request.endpoint not in allowed:
        return redirect(url_for("access", next=request.path))
    return None

@app.context_processor
def inject_access_state():
    return {"access_code_required": access_code_required()}

@app.get("/access")
def access():
    if not access_code_required():
        return redirect(url_for("index"))
    return render_template("access.html", configuration_error=False)

@app.post("/access")
def submit_access_code():
    supplied = request.form.get("access_code", "")
    expected = app.config["DEMO_ACCESS_CODE"]
    if expected and hmac.compare_digest(supplied, expected):
        session["demo_access_granted"] = True
        destination = request.args.get("next", "")
        return redirect(destination if destination.startswith("/") and not destination.startswith("//") else url_for("index"))
    flash("That access code was not accepted.", "error")
    return redirect(url_for("access"))

@app.post("/logout")
def logout():
    session.pop("demo_access_granted", None)
    return redirect(url_for("access"))

def parse_fields(text):
    fields = {}
    for field, patterns in FIELD_PATTERNS.items():
        value = None
        for pattern in patterns:
            match = re.search(pattern, text, flags=re.IGNORECASE | re.MULTILINE)
            if match:
                value = re.sub(r"\s+", " ", match.group(1)).strip(" :\t\r\n")
                if value:
                    break
        fields[field] = value or "Not detected"
    return fields

def extract_docx(path):
    document = Document(path)
    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)

def extract_pdf(path):
    parts = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            page_text = (page.extract_text() or "").strip()
            if page_text:
                parts.append(page_text)
            else:
                try:
                    image = page.to_image(resolution=180).original
                    page_text = pytesseract.image_to_string(image).strip()
                    if page_text:
                        parts.append(page_text)
                except Exception as error:
                    raise ValueError("This PDF appears to be scanned, but OCR could not process it. Install Tesseract OCR and configure TESSERACT_CMD if needed.") from error
    return "\n\n".join(parts)

def extract_image(path):
    try:
        with Image.open(path) as image:
            image.verify()
        with Image.open(path) as image:
            return pytesseract.image_to_string(image).strip()
    except (UnidentifiedImageError, OSError) as error:
        raise ValueError("The image file is invalid or could not be opened.") from error

def validate_file(path, extension):
    if extension == "pdf":
        with open(path, "rb") as uploaded:
            if uploaded.read(5) != b"%PDF-":
                raise ValueError("The selected file is not a valid PDF document.")
    elif extension in {"png", "jpg", "jpeg"}:
        try:
            with Image.open(path) as image:
                image.verify()
                if image.format not in {"PNG", "JPEG"}:
                    raise ValueError("The image contents do not match a supported image type.")
        except (UnidentifiedImageError, OSError) as error:
            raise ValueError("The image file is invalid or could not be opened.") from error
    elif extension == "docx":
        try:
            with zipfile.ZipFile(path) as archive:
                if "word/document.xml" not in archive.namelist():
                    raise ValueError("The DOCX file is missing its document content.")
                if archive.testzip() is not None:
                    raise ValueError("The DOCX file is damaged and could not be read.")
        except zipfile.BadZipFile as error:
            raise ValueError("The selected file is not a valid DOCX document.") from error

@app.get("/")
def index():
    return render_template("index.html")

@app.route("/upload", methods=["GET", "POST"])
def upload():
    if request.method == "GET":
        return render_template("upload.html")
    uploaded_file = request.files.get("certificate")
    if not uploaded_file or not uploaded_file.filename:
        flash("Choose a certificate file to upload.", "error")
        return redirect(url_for("upload"))
    safe_name = secure_filename(uploaded_file.filename)
    if not safe_name or "." not in safe_name:
        flash("The filename is invalid. Rename the file and try again.", "error")
        return redirect(url_for("upload"))
    extension = safe_name.rsplit(".", 1)[-1].lower()
    if extension not in ALLOWED_EXTENSIONS:
        flash("Unsupported file type. Upload a DOCX, PDF, PNG, JPG, or JPEG file.", "error")
        return redirect(url_for("upload"))
    record_id = str(uuid.uuid4())
    storage_dir = Path(app.config["LOCAL_STORAGE"])
    storage_dir.mkdir(parents=True, exist_ok=True)
    stored_path = storage_dir / f"{record_id}.{extension}"
    try:
        with tempfile.TemporaryDirectory(prefix="certificate-") as temporary_directory:
            temporary_path = Path(temporary_directory) / f"{record_id}.{extension}"
            uploaded_file.save(temporary_path)
            file_size = temporary_path.stat().st_size
            if not file_size:
                raise ValueError("The selected file is empty.")
            if file_size > MAX_UPLOAD_BYTES:
                raise ValueError("The file is larger than the 12 MB upload limit.")
            validate_file(temporary_path, extension)
            if extension == "docx":
                extracted_text = extract_docx(temporary_path)
            elif extension == "pdf":
                extracted_text = extract_pdf(temporary_path)
            else:
                extracted_text = extract_image(temporary_path)
            if not extracted_text.strip():
                raise ValueError("No readable text was found. For scanned PDFs or images, install Tesseract OCR.")
            expected_hash = hashlib.sha256(temporary_path.read_bytes()).hexdigest()
            fields = parse_fields(extracted_text)
            stored_path.write_bytes(temporary_path.read_bytes())
        database = get_db()
        database.execute(
            """INSERT INTO certificates
            (id, original_filename, certificate_id, student_name, course_name, issue_date,
             institution, expected_sha256, local_path, uploaded_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (record_id, safe_name, fields["certificate_id"], fields["student_name"],
             fields["course_name"], fields["issue_date"], fields["institution"],
             expected_hash, str(stored_path), datetime.now(timezone.utc).isoformat(timespec="seconds")),
        )
        database.commit()
        return render_template(
            "result.html", filename=safe_name, record_id=record_id,
            expected_hash=expected_hash,
            upload_status="Stored locally for testing (not yet stored in Amazon S3)",
            fields=fields, extracted_text=extracted_text,
        )
    except ValueError as error:
        stored_path.unlink(missing_ok=True)
        flash(str(error), "error")
        return redirect(url_for("upload"))
    except Exception:
        stored_path.unlink(missing_ok=True)
        app.logger.exception("Certificate processing failed")
        flash("The file could not be processed. Check the terminal for the detailed error.", "error")
        return redirect(url_for("upload"))

@app.route("/verify", methods=["GET", "POST"])
def verify():
    database = get_db()
    certificates = database.execute(
        "SELECT id, original_filename, certificate_id, uploaded_at FROM certificates ORDER BY uploaded_at DESC"
    ).fetchall()
    result = None
    if request.method == "POST":
        record_id = request.form.get("record_id", "")
        certificate = database.execute("SELECT * FROM certificates WHERE id = ?", (record_id,)).fetchone()
        if certificate is None:
            flash("Choose a registered certificate from the list.", "error")
        else:
            current_hash = None
            verification_result = "Integrity Mismatch"
            try:
                path = Path(certificate["local_path"])
                if not path.is_file():
                    raise OSError("The locally stored certificate file is missing.")
                digest = hashlib.sha256()
                with path.open("rb") as file:
                    for chunk in iter(lambda: file.read(1024 * 1024), b""):
                        digest.update(chunk)
                current_hash = digest.hexdigest()
                if hmac.compare_digest(current_hash, certificate["expected_sha256"]):
                    verification_result = "Integrity Verified"
            except OSError:
                verification_result = "Verification could not be completed: stored local file is missing."
            database.execute(
                "INSERT INTO verification_events (certificate_ref, verified_at, result) VALUES (?, ?, ?)",
                (record_id, datetime.now(timezone.utc).isoformat(timespec="seconds"), verification_result),
            )
            database.commit()
            result = {
                "certificate": certificate, "current_hash": current_hash or "Not available",
                "expected_hash": certificate["expected_sha256"], "status": verification_result,
                "hashes_match": verification_result == "Integrity Verified",
            }
    return render_template("verify.html", certificates=certificates, result=result)

@app.get("/dashboard")
def dashboard():
    database = get_db()
    total = database.execute("SELECT COUNT(*) AS count FROM certificates").fetchone()["count"]
    successful_checks = database.execute(
        "SELECT COUNT(*) AS count FROM verification_events WHERE result = 'Integrity Verified'"
    ).fetchone()["count"]
    failed_checks = database.execute(
        "SELECT COUNT(*) AS count FROM verification_events WHERE result != 'Integrity Verified'"
    ).fetchone()["count"]
    return render_template("dashboard.html", total=total, successful_uploads=total,
                           successful_checks=successful_checks, failed_checks=failed_checks)

@app.get("/about")
def about():
    return render_template("about.html")

@app.errorhandler(413)
def request_too_large(_error):
    flash("The file is larger than the 12 MB upload limit.", "error")
    return redirect(url_for("upload"))

if __name__ == "__main__":
    app.run(host="127.0.0.1", port=int(os.getenv("PORT", "5000")),
            debug=os.getenv("FLASK_DEBUG") == "1")
