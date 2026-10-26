import hashlib
import io
from pathlib import Path

from docx import Document

import app as certificate_app


class FakeBody(io.BytesIO):
    pass


class FakeS3:
    def __init__(self):
        self.objects = {}
        self.upload_arguments = []

    def upload_file(self, filename, bucket, key, ExtraArgs):
        self.upload_arguments.append((bucket, key, ExtraArgs))
        self.objects[(bucket, key)] = Path(filename).read_bytes()

    def get_object(self, Bucket, Key):
        return {"Body": FakeBody(self.objects[(Bucket, Key)])}


def make_docx():
    document = Document()
    document.add_paragraph("Student Name: Arun Kumar")
    document.add_paragraph("Certificate ID: CERT-2026-001")
    document.add_paragraph("Course Name: Cloud Computing")
    document.add_paragraph("Issue Date: 10 October 2026")
    document.add_paragraph("Issuing Institution: ABC Learning Institute")
    output = io.BytesIO()
    document.save(output)
    output.seek(0)
    return output


def test_upload_registers_hash_and_verification_reads_s3(tmp_path, monkeypatch):
    fake_s3 = FakeS3()
    monkeypatch.setattr(certificate_app, "s3_client", lambda: fake_s3)
    certificate_app.app.config.update(
        TESTING=True,
        DATABASE=str(tmp_path / "certificates.sqlite3"),
        S3_BUCKET="test-private-bucket",
        AWS_REGION="us-east-1",
        DEMO_ACCESS_CODE="",
    )
    with certificate_app.app.app_context():
        certificate_app.init_db()
    client = certificate_app.app.test_client()

    response = client.post(
        "/upload",
        data={"certificate": (make_docx(), "certificate.docx")},
        content_type="multipart/form-data",
    )

    assert response.status_code == 200
    assert b"Arun Kumar" in response.data
    assert b"CERT-2026-001" in response.data
    assert len(fake_s3.objects) == 1
    (bucket, object_key), original_bytes = next(iter(fake_s3.objects.items()))
    expected_hash = hashlib.sha256(original_bytes).hexdigest()
    assert fake_s3.upload_arguments[0][2] == {"ServerSideEncryption": "AES256"}

    database = certificate_app.sqlite3.connect(certificate_app.app.config["DATABASE"])
    record_id, stored_hash = database.execute(
        "SELECT id, expected_sha256 FROM certificates"
    ).fetchone()
    database.close()
    assert stored_hash == expected_hash
    assert bucket == "test-private-bucket"

    response = client.post("/verify", data={"record_id": record_id})
    assert response.status_code == 200
    assert b"Integrity Verified" in response.data
    assert expected_hash.encode() in response.data

    fake_s3.objects[(bucket, object_key)] = original_bytes + b"tampered"
    response = client.post("/verify", data={"record_id": record_id})
    assert response.status_code == 200
    assert b"Integrity Mismatch" in response.data

    response = client.get("/dashboard")
    assert b">1<" in response.data
    assert b"FAILED CHECKS" in response.data


def test_rejects_unsupported_extension_without_s3_call(tmp_path, monkeypatch):
    def unexpected_s3_client():
        raise AssertionError("S3 must not be contacted for an unsupported file")

    monkeypatch.setattr(certificate_app, "s3_client", unexpected_s3_client)
    certificate_app.app.config.update(
        TESTING=True,
        DATABASE=str(tmp_path / "invalid.sqlite3"),
        DEMO_ACCESS_CODE="",
    )
    response = certificate_app.app.test_client().post(
        "/upload",
        data={"certificate": (io.BytesIO(b"nope"), "notes.txt")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert b"Unsupported file type" in response.data


def test_demo_access_code_unlocks_application(tmp_path):
    certificate_app.app.config.update(
        TESTING=True,
        DATABASE=str(tmp_path / "access.sqlite3"),
        DEMO_ACCESS_CODE="classroom-key",
    )
    client = certificate_app.app.test_client()

    response = client.get("/")
    assert response.status_code == 302
    assert "/access" in response.headers["Location"]

    response = client.post("/access", data={"access_code": "classroom-key"})
    assert response.status_code == 302
    assert response.headers["Location"].endswith("/")
    assert client.get("/").status_code == 200
